import re
from io import BytesIO
from pathlib import Path

import fitz
import pdfplumber
import pytesseract
from docx import Document
from PIL import Image


class CVParser:
    @staticmethod
    def read_file(file_path: str) -> str:
        path = Path(file_path)
        if path.suffix.lower() == ".pdf":
            return CVParser._extract_pdf_text(path)
        if path.suffix.lower() == ".docx":
            return CVParser._extract_docx_text(path)
        raise ValueError("Unsupported file type")

    @staticmethod
    def _extract_pdf_text(path: Path) -> str:
        text_chunks = []
        rendered_pdf = None
        with pdfplumber.open(path) as pdf:
            try:
                for page_number, page in enumerate(pdf.pages):
                    text = page.extract_text() or ""
                    text_length = len(text.strip())
                    image_area = sum(
                        max(0, float(image.get("width", 0))) * max(0, float(image.get("height", 0)))
                        for image in page.images
                    )
                    page_area = max(1.0, float(page.width) * float(page.height))
                    image_coverage = min(1.0, image_area / page_area)
                    # Keep native extraction for normal digital pages. Also OCR
                    # pages with a large embedded scan and a thin text layer.
                    if text_length < 80 or (image_coverage >= 0.35 and text_length < 800):
                        if rendered_pdf is None:
                            rendered_pdf = fitz.open(path)
                        ocr_text = CVParser._ocr_pdf_page(rendered_pdf, page_number)
                        if len(ocr_text.strip()) > text_length:
                            text = ocr_text
                    text_chunks.append(text)
            finally:
                if rendered_pdf is not None:
                    rendered_pdf.close()
        return "\n".join(text_chunks)

    @staticmethod
    def _ocr_pdf_page(document, page_number: int) -> str:
        page = document.load_page(page_number)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        with Image.open(BytesIO(pixmap.tobytes("png"))) as image:
            return pytesseract.image_to_string(image, lang="eng+vie", config="--psm 3")

    @staticmethod
    def _extract_docx_text(path: Path) -> str:
        document = Document(path)
        chunks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    chunks.append(" | ".join(cells))
        for section in document.sections:
            for part in (section.header, section.footer):
                chunks.extend(paragraph.text.strip() for paragraph in part.paragraphs if paragraph.text.strip())
        return "\n".join(chunks)

    @staticmethod
    def extract_contact_details(text: str):
        email = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)
        phone = re.search(r"(?:\+?\d[\d\s().-]{7,}\d)", text)
        return {
            "email": email.group(0) if email else None,
            "phone": phone.group(0).strip() if phone else None,
        }

    @staticmethod
    def clean_text(text: str) -> str:
        cleaned = re.sub(r"\s+", " ", text or "")
        return cleaned.strip()
