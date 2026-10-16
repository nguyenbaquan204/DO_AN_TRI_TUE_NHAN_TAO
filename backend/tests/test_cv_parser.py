from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document

from app.services.cv_parser import CVParser


def _mock_pdf_page(text: str, images=None):
    page = MagicMock()
    page.extract_text.return_value = text
    page.images = images or []
    page.width = 100
    page.height = 100
    return page


def _open_mock(value):
    value.__enter__.return_value = value
    value.__exit__.return_value = False
    return value


def test_pdf_parser_uses_native_text_without_ocr():
    page = _mock_pdf_page("A digital resume with enough searchable text. " * 4)
    pdf = _open_mock(MagicMock(pages=[page]))
    rendered = _open_mock(MagicMock())
    with patch("app.services.cv_parser.pdfplumber.open", return_value=pdf), \
            patch("app.services.cv_parser.fitz.open", return_value=rendered) as open_rendered_pdf, \
            patch.object(CVParser, "_ocr_pdf_page") as ocr:
        text = CVParser._extract_pdf_text(Path("digital.pdf"))
    assert text.startswith("A digital resume")
    open_rendered_pdf.assert_not_called()
    ocr.assert_not_called()


def test_pdf_parser_ocr_fallback_for_scanned_page():
    page = _mock_pdf_page("")
    pdf = _open_mock(MagicMock(pages=[page]))
    rendered = _open_mock(MagicMock())
    with patch("app.services.cv_parser.pdfplumber.open", return_value=pdf), \
            patch("app.services.cv_parser.fitz.open", return_value=rendered), \
            patch.object(CVParser, "_ocr_pdf_page", return_value="Scanned CV content") as ocr:
        text = CVParser._extract_pdf_text(Path("scanned.pdf"))
    assert text == "Scanned CV content"
    ocr.assert_called_once_with(rendered, 0)


def test_pdf_parser_ocr_fallback_for_image_heavy_mixed_page():
    page = _mock_pdf_page("short header", images=[{"width": 90, "height": 90}])
    pdf = _open_mock(MagicMock(pages=[page]))
    rendered = _open_mock(MagicMock())
    with patch("app.services.cv_parser.pdfplumber.open", return_value=pdf), \
            patch("app.services.cv_parser.fitz.open", return_value=rendered), \
            patch.object(CVParser, "_ocr_pdf_page", return_value="Long OCR body with all scanned page text"):
        text = CVParser._extract_pdf_text(Path("mixed.pdf"))
    assert text == "Long OCR body with all scanned page text"


def test_docx_parser_reads_table_content(tmp_path):
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "Backend Engineer"
    path = tmp_path / "table-cv.docx"
    document.save(path)
    text = CVParser._extract_docx_text(path)
    assert "Python | Backend Engineer" in text
